"""Shared Deep Agents construction logic for the Career Trajectory Optimizer.

This mirrors the exact technical pattern used in `career_optimizer_deep_agent.ipynb`
(deep agents, specialist subagents, composite backend, store, checkpointer, interrupts),
factored out so the Streamlit UI (`app.py`) doesn't duplicate it.
"""
from __future__ import annotations

import os
from pathlib import Path

from langchain.tools import tool
from deepagents import create_deep_agent
from deepagents.backends import CompositeBackend, FilesystemBackend, StoreBackend
from langgraph.checkpoint.memory import MemorySaver
from langgraph.store.memory import InMemoryStore

PROJECT_DIR = Path(__file__).resolve().parent
GROUNDING_PATH = PROJECT_DIR / "grounding" / "india_market_context.md"
SKILLS = ["/skills/"]


def load_grounding() -> str:
    return GROUNDING_PATH.read_text()


RESUME_FILE_TYPES = ["pdf", "docx", "txt"]


def extract_resume_text(filename: str, file_bytes: bytes) -> str:
    """Extract raw text from an uploaded resume (.pdf, .docx, or .txt).

    The extracted text is still unstructured — the `profile-analyzer` subagent does
    the actual field extraction (role, company, skills, education, etc.).
    """
    import io

    suffix = Path(filename).suffix.lower().lstrip(".")
    if suffix == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(file_bytes))
        return "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    if suffix == "docx":
        return _extract_docx_text(file_bytes)
    if suffix == "txt":
        return file_bytes.decode("utf-8", errors="ignore").strip()
    raise ValueError(f"Unsupported resume file type '.{suffix}' — use .pdf, .docx, or .txt.")


def _extract_docx_text(file_bytes: bytes) -> str:
    """Walk the whole document — many resumes lay out skills/experience in tables,
    which plain `doc.paragraphs` iteration silently skips."""
    import io
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(file_bytes))
    lines: list[str] = []

    def iter_block_items(parent):
        parent_elm = parent.element.body if hasattr(parent, "element") else parent._tc
        for child in parent_elm.iterchildren():
            if child.tag == qn("w:p"):
                yield Paragraph(child, parent)
            elif child.tag == qn("w:tbl"):
                yield Table(child, parent)

    def walk(parent):
        for block in iter_block_items(parent):
            if isinstance(block, Paragraph):
                if block.text.strip():
                    lines.append(block.text)
            elif isinstance(block, Table):
                for row in block.rows:
                    for cell in row.cells:
                        walk(cell)  # recurse for nested tables/paragraphs

    walk(doc)

    for section in doc.sections:
        for p in list(section.header.paragraphs) + list(section.footer.paragraphs):
            if p.text.strip():
                lines.append(p.text)

    return "\n".join(lines).strip()


_PLACEHOLDER = "your-key-here"


def _is_real_key(value: str) -> bool:
    """Guards against unfilled `.env.example`-style placeholder values."""
    v = value.strip()
    return bool(v) and v.lower() != _PLACEHOLDER


def build_research_tools() -> list:
    """Web search is optional. Prefer LinkUp as the primary provider; fall back to
    Tavily if LinkUp isn't configured or fails to initialize; otherwise no search
    tool is added (subagents fall back to grounding-only)."""
    if _is_real_key(os.environ.get("LINKUP_API_KEY", "")):
        try:
            from langchain_linkup import LinkupSearchTool
            return [LinkupSearchTool(depth="standard", output_type="searchResults", max_results=3)]
        except Exception:
            pass  # fall through to Tavily below
    if _is_real_key(os.environ.get("TAVILY_API_KEY", "")):
        from langchain_tavily import TavilySearch
        return [TavilySearch(max_results=3)]
    return []


@tool
def roadmap_validator(text: str) -> str:
    """Validate a career roadmap (required sections present). Returns a pass/fail report."""
    required = ["goal", "gap", "upskill", "30", "60", "90", "timeline", "salary", "probability"]
    lowered = text.lower()
    missing = [word for word in required if word not in lowered]
    return "PASS" if not missing else "ISSUES:\n- Missing sections/keywords: " + ", ".join(missing)


@tool
def risk_dashboard_validator(salary_risk: int, skill_risk: int, security_risk: int,
                              burnout_risk: int, growth_risk: int) -> str:
    """Validate a risk dashboard: all five risk scores must be present and within 0-10."""
    scores = {
        "salary_risk": salary_risk,
        "skill_risk": skill_risk,
        "security_risk": security_risk,
        "burnout_risk": burnout_risk,
        "growth_risk": growth_risk,
    }
    issues = [f"{name}={value} out of range 0-10" for name, value in scores.items() if not (0 <= value <= 10)]
    return "PASS" if not issues else "ISSUES:\n- " + "\n- ".join(issues)


@tool
def sync_to_calendar(task_description: str, cadence: str) -> str:
    """Sync a roadmap task to a calendar (SIMULATED — nothing is actually synced). cadence: 'daily', 'weekly', 'monthly', or 'quarterly'."""
    return f"[SIMULATED] Synced '{task_description}' ({cadence}) to calendar. Nothing was actually synced."


def build_subagents(research_tools: list) -> list[dict]:
    profile_analyzer_subagent = {
        "name": "profile-analyzer",
        "description": "Parses the raw profile/resume text into a structured profile and builds a career timeline. Use first, always.",
        "system_prompt": (
            "You parse a raw career profile (resume text or manual entry). Extract role, company, "
            "years of experience, skills, and education. Detect education tier (IIT/NIT = premium "
            "tier per grounding). Build a timeline of past roles and flag job-hopping or long tenure. "
            "Write the structured profile and timeline to /profile/timeline.md and return a one-line summary."
        ),
        "tools": [],
    }

    market_researcher_subagent = {
        "name": "market-researcher",
        "description": "Gathers live Indian tech market data (salary trends, hiring demand, GenAI growth) via web search. Use second, right after profile-analyzer.",
        "system_prompt": (
            "Research the current Indian tech job market using web_search if available: salary "
            "benchmarks for the candidate's role/experience/location, hiring demand for their "
            "skills, and GenAI/emerging-skill growth trends. If no web_search tool is available, "
            "rely on the grounding facts only and say so explicitly. Write 6-10 crisp, sourced "
            "bullet points to /research/market_brief.md and return a one-line summary."
        ),
        "tools": research_tools,
    }

    salary_intelligence_subagent = {
        "name": "salary-intelligence",
        "description": "Benchmarks the candidate's current pay against the Indian market. Use after market-researcher.",
        "system_prompt": (
            "Load the 'salary-benchmarking' skill and follow it strictly. Read /profile/timeline.md "
            "and /research/market_brief.md. Ground every number in those two files only — do not "
            "search the web yourself. Produce a market range, the candidate's estimated position in "
            "it, and (if a goal salary was given in the brief) the gap and levers to close it. "
            "Save the result to /analysis/salary.md and return a one-line summary."
        ),
        "tools": [],
        "skills": SKILLS,
    }

    market_tracker_subagent = {
        "name": "market-tracker",
        "description": "Tracks skill demand and job-market trends relevant to the candidate's skills. Use after market-researcher.",
        "system_prompt": (
            "Read /profile/timeline.md and /research/market_brief.md. Do not search the web "
            "yourself — reason only from those two files plus the grounding facts (e.g. GenAI "
            "~180% YoY growth). Identify which of the candidate's skills are growing, flat, or "
            "declining, and which emerging skills are worth learning. Save to "
            "/analysis/market_trends.md and return a one-line summary."
        ),
        "tools": [],
    }

    risk_scorer_subagent = {
        "name": "risk-scorer",
        "description": "Scores readiness for the next level and produces a 5-axis risk dashboard. Use after salary-intelligence and market-tracker.",
        "system_prompt": (
            "Read /profile/timeline.md, /analysis/salary.md, and /analysis/market_trends.md. Score "
            "readiness for the next career level from the candidate's accomplishments. Then score "
            "five risks 0-10 each: salary_risk, skill_risk, security_risk, burnout_risk, growth_risk. "
            "Run risk_dashboard_validator on your five scores and fix any issues it reports. Save the "
            "full dashboard to /analysis/risk_dashboard.md and return a one-line summary."
        ),
        "tools": [risk_dashboard_validator],
    }

    career_recommender_subagent = {
        "name": "career-recommender",
        "description": (
            "Builds goal-based recommendations for Career Aspiration Mode only (a stated salary, "
            "company, or role goal). Use after risk-scorer, only when the user gave a goal."
        ),
        "system_prompt": (
            "Read /profile/timeline.md, /research/market_brief.md, /analysis/salary.md, "
            "/analysis/market_trends.md, and /analysis/risk_dashboard.md. Do not search the web "
            "yourself — use only those files. The user's goal is embedded in the task you were "
            "given. Apply this logic: goal salary -> identify target company tier and required "
            "upskills; goal company (e.g. Microsoft) -> use /research/market_brief.md for the "
            "company's bar and gap vs. candidate; goal role (e.g. Tech Lead) -> identify "
            "leadership/architecture gaps and visibility actions; goal 'GenAI Engineer' -> note the "
            "~20-30% skill premium and a 3-6 month learning plan; goal US remote/relocation -> note "
            "the ~20-30% feasibility and 12-18 month prep. Save the gap analysis and recommendation "
            "to /analysis/recommendation.md and return a one-line summary."
        ),
        "tools": [],
    }

    roadmap_planner_subagent = {
        "name": "roadmap-planner",
        "description": "Turns the recommendation into a 30-60-90 day roadmap. Use only in Career Aspiration Mode, after career-recommender.",
        "system_prompt": (
            "Load the 'roadmap-planning' skill and follow it strictly. Read /analysis/recommendation.md. "
            "Produce: goal summary, gap analysis, upskilling plan, a 30-60-90 day plan, timeline to "
            "goal, salary projection, and a success probability percentage. Run roadmap_validator on "
            "your draft and fix any issues it reports. Save the roadmap to /roadmap/plan.md and return "
            "the final roadmap text."
        ),
        "tools": [roadmap_validator],
        "skills": SKILLS,
    }

    critic_subagent = {
        "name": "critic",
        "description": (
            "Reviews all analysis files for grounding accuracy, methodology compliance, and quality; "
            "runs the validators; suggests fixes."
        ),
        "system_prompt": (
            "Read every file that exists under /profile, /research, /analysis, and /roadmap. Review "
            "on four axes:\n"
            "1. GROUNDING ACCURACY — flag any number or claim not traceable to the grounding file or "
            "/research/market_brief.md.\n"
            "2. METHODOLOGY COMPLIANCE — load the 'salary-benchmarking' and 'roadmap-planning' skills "
            "and check the relevant files actually follow the methodology.\n"
            "3. INTERNAL CONSISTENCY — do the salary, risk, and roadmap files agree with each other?\n"
            "4. VALIDATORS — run roadmap_validator (if a roadmap exists) and risk_dashboard_validator "
            "(if a risk dashboard exists).\n"
            "Return a short review with concrete, actionable fixes, calling out any unsupported "
            "numbers or skill violations explicitly."
        ),
        "tools": [roadmap_validator, risk_dashboard_validator],
        "skills": SKILLS,
    }

    return [
        profile_analyzer_subagent,
        market_researcher_subagent,
        salary_intelligence_subagent,
        market_tracker_subagent,
        risk_scorer_subagent,
        career_recommender_subagent,
        roadmap_planner_subagent,
        critic_subagent,
    ]


def build_orchestrator_prompt(grounding: str) -> str:
    return f"""You are an AI career coach for Indian tech professionals.
For a profile you produce either a market-position analysis or a full goal-based roadmap.

Market grounding (always honor this; never invent facts beyond it + research):
{grounding}

Workflow — first call write_todos to plan, then:
1. Delegate to 'profile-analyzer' to parse the profile (writes /profile/timeline.md).
2. Delegate to 'market-researcher' to gather live Indian tech market data via web search
   (writes /research/market_brief.md).
3. Delegate to 'salary-intelligence' AND 'market-tracker' CONCURRENTLY — emit both `task`
   tool calls in the SAME step (they are independent: both only need profile-analyzer's
   and market-researcher's output, not each other's). This halves wall-clock time versus
   running them one after another.
4. Delegate to 'risk-scorer' (writes /analysis/risk_dashboard.md) — needs both files from
   step 3, so this must wait until both have returned.
5. INTENT DETECTION: read the user's brief.
   - If it states a concrete goal (a target salary, company, or role/timeline), this is
     CAREER ASPIRATION MODE: delegate to 'career-recommender' (writes
     /analysis/recommendation.md), then 'roadmap-planner' (writes /roadmap/plan.md).
   - Otherwise this is PROFILE ANALYSIS MODE: skip step 5's subagents and go straight to
     step 6 with only the salary/market/risk analysis.
6. Delegate to 'critic' to review everything produced so far; if it flags something
   important, re-delegate to fix it. Tell the critic to read all relevant files in as few
   read_file calls as possible to keep this step fast.
7. Save a one-paragraph progress note (mode used, key findings, current goal if any) to
   /memories/career_progress.md for future coaching sessions.
8. ONLY in Career Aspiration Mode: sync the roadmap's monthly checkpoints to a calendar.
   Emit one sync_to_calendar call per checkpoint (cadence='monthly' for month-end reviews,
   cadence='weekly' for milestones). Do not write a final summary until all sync_to_calendar
   calls have been made and returned. In Profile Analysis Mode, skip this step entirely.

Keep everything concise, data-grounded, and honest about timelines."""


def build_agent(model: str = "gpt-5-mini"):
    """Assemble the deep agent. Returns (agent, backend, store, checkpointer)."""
    grounding = load_grounding()
    research_tools = build_research_tools()
    subagents = build_subagents(research_tools)

    store = InMemoryStore()        # cross-session memory (use a PostgresStore in production)
    checkpointer = MemorySaver()   # per-thread state + required for HITL interrupts

    backend = CompositeBackend(
        default=FilesystemBackend(root_dir=str(PROJECT_DIR), virtual_mode=True),
        routes={"/memories/": StoreBackend(namespace=lambda ctx: ("career",))},
    )

    # Bound retries/timeout explicitly. Left at library defaults, a 429 (rate
    # limit) can silently back off for many minutes on a free-tier daily quota
    # with tight per-minute limits — fail fast with a visible error instead.
    from langchain_openai import ChatOpenAI
    llm = ChatOpenAI(model=model, max_retries=2, timeout=60)

    agent = create_deep_agent(
        model=llm,
        system_prompt=build_orchestrator_prompt(grounding),
        tools=[sync_to_calendar],
        subagents=subagents,
        backend=backend,
        store=store,
        skills=SKILLS,
        interrupt_on={"sync_to_calendar": True},   # HITL gate — needs the checkpointer
        checkpointer=checkpointer,
    )
    return agent, backend, store, checkpointer


def pending_requests(result: dict) -> list:
    """Unwrap the pending HITL action requests from an agent.invoke() result."""
    reqs: list = []
    for intr in result.get("__interrupt__", []):
        val = intr.value
        if isinstance(val, dict) and "action_requests" in val:
            reqs.extend(val["action_requests"])
        elif isinstance(val, list):
            reqs.extend(val)
        else:
            reqs.append(val)
    return reqs


def _format_todos(todos: list) -> str:
    lines = []
    for t in todos or []:
        status = t.get("status", "pending") if isinstance(t, dict) else "pending"
        content = t.get("content", str(t)) if isinstance(t, dict) else str(t)
        mark = {"completed": "[x]", "in_progress": "[~]"}.get(status, "[ ]")
        lines.append(f"- {mark} {content}")
    return "\n".join(lines) if lines else "(no todos)"


def build_trace_entries(messages: list) -> list[dict]:
    """Parse a raw `result["messages"]` list into a sequential, labeled trace so
    each stage (planning, delegation, subagent result, final answer) is clearly
    headed and traceable, instead of a flat dump of role + content pairs."""
    entries: list[dict] = []
    pending_calls: dict[str, dict] = {}  # tool_call_id -> {"name": ..., "args": ...}

    for m in messages:
        msg_type = getattr(m, "type", None)

        if msg_type == "human":
            entries.append({"kind": "user", "heading": "User request", "body": m.content})
            continue

        if msg_type == "ai":
            for tc in (getattr(m, "tool_calls", None) or []):
                name = tc.get("name", "tool")
                args = tc.get("args", {}) or {}
                call_id = tc.get("id")
                pending_calls[call_id] = {"name": name, "args": args}

                if name == "task":
                    subagent = args.get("subagent_type", "subagent")
                    entries.append({
                        "kind": "delegate",
                        "heading": f"Orchestrator delegates → `{subagent}`",
                        "body": args.get("description", ""),
                    })
                elif name == "write_todos":
                    entries.append({
                        "kind": "plan",
                        "heading": "Orchestrator plans (write_todos)",
                        "body": _format_todos(args.get("todos", [])),
                    })
                elif name == "sync_to_calendar":
                    entries.append({
                        "kind": "tool_call",
                        "heading": "Orchestrator requests calendar sync",
                        "body": f"{args.get('task_description', '')} (cadence: {args.get('cadence', '?')})",
                    })
                else:
                    entries.append({
                        "kind": "tool_call",
                        "heading": f"Tool call → `{name}`",
                        "body": str(args),
                    })

            content = getattr(m, "content", None)
            if content:
                entries.append({"kind": "final", "heading": "Agent response", "body": content})
            continue

        if msg_type == "tool":
            call_id = getattr(m, "tool_call_id", None)
            info = pending_calls.get(call_id, {})
            name = getattr(m, "name", None) or info.get("name", "tool")
            if name == "task":
                subagent = info.get("args", {}).get("subagent_type", "subagent")
                heading = f"Result from `{subagent}`"
            else:
                heading = f"Result of `{name}`"
            entries.append({"kind": "tool_result", "heading": heading, "body": m.content})
            continue

    return entries
